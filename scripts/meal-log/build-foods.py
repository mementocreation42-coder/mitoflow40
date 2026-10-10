"""日本食品標準成分表（八訂）増補2023年の「第2章（データ）」Excel から、
食事記録（/meal-log）で使う食品データ public/data/foods-mext.json を作る。

使い方: python3 scripts/meal-log/build-foods.py <ダウンロードした .xlsx>
出典: https://www.mext.go.jp/a_menu/syokuhinseibun/mext_00001.html
"""
import json
import re
import sys

import openpyxl

# 成分識別子 → アプリ内のキー（lib/nutrition.ts の NUTRIENTS と対応）
FIELDS = {
    'ENERC_KCAL': 'kcal', 'PROT-': 'prot', 'FAT-': 'fat', 'CHOCDF-': 'carb', 'CHOAVLDF-': 'avcarb',
    'FIB-': 'fib', 'NACL_EQ': 'salt', 'CHOLE': 'chole',
    'VITA_RAE': 'vita', 'VITD': 'vitd', 'TOCPHA': 'vite', 'VITK': 'vitk', 'THIA': 'b1', 'RIBF': 'b2',
    'NE': 'ne', 'VITB6A': 'b6', 'VITB12': 'b12', 'FOL': 'fol', 'PANTAC': 'pant', 'BIOT': 'biot', 'VITC': 'vitc',
    'K': 'k', 'CA': 'ca', 'MG': 'mg', 'P': 'p', 'FE': 'fe', 'ZN': 'zn', 'CU': 'cu', 'MN': 'mn',
    'ID': 'id', 'SE': 'se', 'CR': 'cr', 'MO': 'mo', 'NA': 'na',
}


def value(cell):
    """数値・推計値 (x)・Tr（微量）→ 数値、- や空欄（未測定）→ None。"""
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return cell
    s = str(cell).strip().strip('()')
    if s in ('', '-', '*'):
        return None
    if s == 'Tr':
        return 0
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['表全体']
    rows = list(ws.iter_rows(values_only=True))
    updated = next((c for c in rows[0] if isinstance(c, str) and '更新日' in c), '')
    ids = [str(c).strip() if c else None for c in rows[11]]
    keys = list(FIELDS.values())
    col = {FIELDS[i]: n for n, i in enumerate(ids) if i in FIELDS}
    missing = set(keys) - set(col)
    assert not missing, f'見つからない成分: {missing}'

    groups = {}
    for name in wb.sheetnames[1:]:
        m = re.match(r'(\d+)(.+)', name)
        groups[m.group(1).zfill(2)] = m.group(2)

    foods = []
    for r in rows[12:]:
        if not r[1] or not r[3]:
            continue
        name = re.sub(r'[　 ]+', ' ', str(r[3])).strip()
        foods.append([str(r[1]), name] + [value(r[col[k]]) for k in keys])

    out = {
        'source': '日本食品標準成分表（八訂）増補2023年（文部科学省）',
        'updated': updated.replace('更新日：', ''),
        'per': '可食部100g当たり',
        'keys': keys,
        'groups': groups,
        'foods': foods,
    }
    with open('public/data/foods-mext.json', 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
    print(len(foods), 'foods')


if __name__ == '__main__':
    main(sys.argv[1])
